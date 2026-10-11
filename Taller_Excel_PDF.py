from openpyxl import load_workbook
from tabulate import tabulate
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

# Cargar los datos de Excel
wb = load_workbook('ventas_tecnologia_v2.xlsx', data_only=True)
sheet = wb.active
headers = [cell.value for cell in sheet[1]]

data = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    data.append(list(row))

alignments = ("center", "left", "center", "right", "right")

# Configurar los estilos del PDF
styles = getSampleStyleSheet()
title_style = ParagraphStyle(
    "TitleStyle", parent=styles["Title"], fontName="Helvetica-Bold",
    fontSize=18, alignment=TA_CENTER, textColor=colors.HexColor("#17365D")
)
body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontSize=9)
header_style = ParagraphStyle(
    "HeaderStyle", parent=styles["Normal"], fontName="Helvetica-Bold",
    fontSize=10, textColor=colors.whitesmoke, alignment=1
)

# Mostrar reporte de ventas en consola
def mostrar_catalogo():
    print(f"\n==================== REPORTE DE VENTAS EN CONSOLA ({len(data)} PRODUCTOS) ====================")
    print(tabulate(data, headers=headers, tablefmt="fancy_grid", colalign=alignments))

# Reto 1: Filtrar productos con precio superior a $100
def filtrar_productos():
    filtrados = []
    for producto in data:
        if producto[3] > 100:
            filtrados.append(producto)

    print("\nPRODUCTOS CON PRECIO SUPERIOR A $100")
    print(tabulate(filtrados, headers=headers, tablefmt="fancy_grid", colalign=alignments))
    exportar_pdf(filtrados, "Productos con precio superior a $100")

# Reto 2: Calcular unidades vendidas y precio promedio
def mostrar_resumen():
    total_unidades = 0
    suma_precios = 0
    for producto in data:
        total_unidades += producto[4]
        suma_precios += producto[3]

    if data:
        precio_promedio = suma_precios / len(data)
    else:
        precio_promedio = 0

    print(f"\nTotal de unidades vendidas: {total_unidades}")
    print(f"Precio promedio: ${precio_promedio:,.2f}")

    # Agregar una fila de resumen al PDF
    resumen = ["", "RESUMEN", "", f"Promedio: ${precio_promedio:,.2f}", total_unidades]
    exportar_pdf(data, "Informe Ejecutivo de Ventas", resumen)

# Reto 3: Buscar un producto por ID
def buscar_producto():
    try:
        id_buscar = int(input("Ingrese el ID del producto: "))
        producto = next((r for r in data if r[0] == id_buscar), None)

        if producto:
            print(tabulate([producto], headers=headers, tablefmt="fancy_grid", colalign=alignments))
        else:
            print("No se encontró un producto con ese ID.")
    except ValueError:
        print("Ingrese un ID numérico válido.")

# Crear el PDF
def exportar_pdf(registros, titulo, resumen=None):
    doc = SimpleDocTemplate("reporte_ejecutivo_ventas.pdf", pagesize=letter)
    story = [
        Paragraph(titulo, title_style),
        Paragraph(f"<b>Total de productos procesados:</b> {len(registros)}", body_style),
        Spacer(1, 12)
    ]

    # Preparar encabezados y filas
    tabla_datos = [[Paragraph(str(h), header_style) for h in headers]]

    for r in registros:
        tabla_datos.append([
            Paragraph(str(r[0]), body_style),
            Paragraph(str(r[1]), body_style),
            Paragraph(str(r[2]), body_style),
            Paragraph(f"${r[3]:,.2f}", body_style),
            Paragraph(str(r[4]), body_style)
        ])

    # Agregar resumen
    if resumen is not None:
        tabla_datos.append([Paragraph(str(c), body_style) for c in resumen])

    pdf_table = Table(tabla_datos, colWidths=[40, 160, 100, 80, 80], repeatRows=1)
    pdf_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey])
    ]))

    story.append(pdf_table)
    doc.build(story)
    print("\n[ÉXITO] PDF generado: reporte_ejecutivo_ventas.pdf")

# Reto 4: Menú interactivo
while True:
    print("\n===== MENÚ TECNOVENTAS =====")
    print("1. Mostrar catálogo completo")
    print("2. Buscar producto por ID")
    print("3. Filtrar productos con precio superior a $100")
    print("4. Mostrar resumen de ventas y exportar PDF")
    print("0. Salir")
    opcion = input("Seleccione una opción: ")

    if opcion == "1":
        mostrar_catalogo()
    elif opcion == "2":
        buscar_producto()
    elif opcion == "3":
        filtrar_productos()
    elif opcion == "4":
        mostrar_resumen()
    elif opcion == "0":
        print("¡Hasta luego!")
        break
    else:
        print("Opción no válida. Intente nuevamente.")